import io
import json
import re
from typing import Any, Dict, List

import streamlit as st
from docx import Document
from google import genai
from pypdf import PdfReader


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="AI Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

MODEL_NAME = "gemini-3.5-flash"


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>
    .main-title {
        font-size: 42px;
        font-weight: 800;
        margin-bottom: 0;
    }

    .subtitle {
        font-size: 17px;
        opacity: 0.75;
        margin-bottom: 25px;
    }

    .score-card {
        padding: 24px;
        border: 1px solid rgba(128,128,128,.25);
        border-radius: 16px;
        text-align: center;
        margin-bottom: 15px;
    }

    .score-number {
        font-size: 54px;
        font-weight: 800;
    }

    .score-label {
        font-size: 18px;
        font-weight: 600;
    }

    .section-card {
        padding: 18px;
        border: 1px solid rgba(128,128,128,.20);
        border-radius: 14px;
        margin-bottom: 12px;
    }

    .small-muted {
        font-size: 13px;
        opacity: 0.65;
    }

    div[data-testid="stMetricValue"] {
        font-size: 30px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SESSION STATE
# ============================================================

if "analysis" not in st.session_state:
    st.session_state.analysis = None

if "resume_text" not in st.session_state:
    st.session_state.resume_text = ""

if "resume_name" not in st.session_state:
    st.session_state.resume_name = ""


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []

    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")

    return "\n".join(pages).strip()


def extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts: List[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore").strip()


def extract_resume(uploaded_file) -> str:
    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]
    data = uploaded_file.getvalue()

    if extension == "pdf":
        return extract_pdf(data)

    if extension == "docx":
        return extract_docx(data)

    if extension == "txt":
        return extract_txt(data)

    raise ValueError("Unsupported file type.")


# ============================================================
# TEXT / ATS HELPERS
# ============================================================

def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower()).strip()


def count_words(text: str) -> int:
    return len(re.findall(r"\b[\w+#.-]+\b", text))


def extract_email(text: str) -> bool:
    return bool(
        re.search(
            r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
            text,
            re.I,
        )
    )


def extract_phone(text: str) -> bool:
    return bool(
        re.search(
            r"(?<!\d)(?:\+?\d[\d\s().-]{7,}\d)(?!\d)",
            text,
        )
    )


def has_section(text: str, patterns: List[str]) -> bool:
    lowered = normalize_text(text)
    return any(re.search(pattern, lowered) for pattern in patterns)


def calculate_local_ats(
    resume_text: str,
    job_description: str = "",
) -> Dict[str, Any]:
    text = normalize_text(resume_text)
    words = count_words(resume_text)

    checks = {
        "Contact information": extract_email(resume_text)
        or extract_phone(resume_text),

        "Professional summary": has_section(
            resume_text,
            [
                r"\bprofessional summary\b",
                r"\bsummary\b",
                r"\bprofile\b",
                r"\bobjective\b",
            ],
        ),

        "Work experience": has_section(
            resume_text,
            [
                r"\bwork experience\b",
                r"\bprofessional experience\b",
                r"\bexperience\b",
                r"\bemployment history\b",
            ],
        ),

        "Education": has_section(
            resume_text,
            [
                r"\beducation\b",
                r"\bacademic background\b",
                r"\bqualifications\b",
            ],
        ),

        "Skills": has_section(
            resume_text,
            [
                r"\btechnical skills\b",
                r"\bskills\b",
                r"\bcore competencies\b",
            ],
        ),

        "Action verbs": len(
            re.findall(
                r"\b("
                r"achieved|analyzed|built|created|designed|"
                r"developed|delivered|implemented|improved|"
                r"increased|led|managed|optimized|reduced|"
                r"automated|launched|maintained|resolved|"
                r"coordinated|generated"
                r")\b",
                text,
            )
        ) >= 3,

        "Quantified achievements": bool(
            re.search(
                r"\b\d+(?:\.\d+)?\s*(?:%|percent|k|m|million|"
                r"thousand|years?|months?|users?|clients?|projects?)\b",
                text,
                re.I,
            )
        ),

        "Reasonable resume length": 250 <= words <= 1400,
    }

    structure_score = round(
        sum(checks.values()) / len(checks) * 100
    )

    keyword_result = {
        "score": None,
        "matched": [],
        "missing": [],
        "total_keywords": 0,
    }

    if job_description.strip():
        stop_words = {
            "the", "and", "for", "with", "that", "this", "from",
            "your", "you", "our", "are", "will", "have", "has",
            "job", "role", "work", "years", "year", "into",
            "their", "they", "about", "what", "who", "can",
            "should", "must", "need", "using", "used", "use",
            "including", "such", "also", "more", "than", "all",
            "other", "some", "responsible", "required",
            "preferred", "candidate", "position", "company",
        }

        jd_terms = re.findall(
            r"[a-zA-Z][a-zA-Z0-9+#./-]{2,}",
            normalize_text(job_description),
        )

        keywords = []
        seen = set()

        for word in jd_terms:
            if word not in stop_words and word not in seen:
                keywords.append(word)
                seen.add(word)

        matched = []
        missing = []

        for keyword in keywords:
            pattern = r"(?<!\w)" + re.escape(keyword) + r"(?!\w)"
            if re.search(pattern, text):
                matched.append(keyword)
            else:
                missing.append(keyword)

        keyword_score = round(
            len(matched) / max(len(keywords), 1) * 100
        )

        keyword_result = {
            "score": keyword_score,
            "matched": matched[:100],
            "missing": missing[:100],
            "total_keywords": len(keywords),
        }

    if keyword_result["score"] is not None:
        final_score = round(
            structure_score * 0.55
            + keyword_result["score"] * 0.45
        )
    else:
        final_score = structure_score

    return {
        "score": max(0, min(100, final_score)),
        "structure_score": structure_score,
        "checks": checks,
        "word_count": words,
        "keyword": keyword_result,
    }


# ============================================================
# GEMINI
# ============================================================

def get_api_key() -> str:
    try:
        secret_key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        secret_key = ""

    if secret_key:
        return str(secret_key).strip()

    return str(
        st.session_state.get("api_key", "")
    ).strip()


def clean_json_response(raw_text: str) -> str:
    cleaned = raw_text.strip()

    cleaned = re.sub(
        r"^```json\s*",
        "",
        cleaned,
        flags=re.I,
    )

    cleaned = re.sub(
        r"^```\s*",
        "",
        cleaned,
    )

    cleaned = re.sub(
        r"\s*```$",
        "",
        cleaned,
    )

    start = cleaned.find("{")
    end = cleaned.rfind("}")

    if start != -1 and end != -1:
        cleaned = cleaned[start:end + 1]

    return cleaned.strip()


def gemini_analyze(
    resume_text: str,
    job_description: str,
    api_key: str,
) -> Dict[str, Any]:

    client = genai.Client(api_key=api_key)

    prompt = f"""
You are an expert Applicant Tracking System (ATS) resume
reviewer, recruiter, and career coach.

Analyze the resume carefully.

If a job description is provided, evaluate how well the
resume matches that specific job.

Return ONLY valid JSON.
Do not use markdown.
Do not use code fences.
Do not add text before or after the JSON.

Use exactly this structure:

{{
  "ats_score": 0,
  "summary": "",
  "score_breakdown": {{
    "keyword_match": 0,
    "ats_formatting": 0,
    "skills_relevance": 0,
    "experience_relevance": 0,
    "achievements": 0,
    "resume_structure": 0
  }},
  "strengths": [],
  "top_improvements": [],
  "missing_keywords": [],
  "matched_keywords": [],
  "section_feedback": {{
    "Contact": "",
    "Professional Summary": "",
    "Work Experience": "",
    "Education": "",
    "Skills": "",
    "Certifications": "",
    "Projects": ""
  }},
  "rewrites": [
    {{
      "before": "",
      "after": "",
      "reason": ""
    }}
  ],
  "formatting_tips": [],
  "recommended_skills": [],
  "ats_warnings": []
}}

Important rules:

1. ats_score must be an integer from 0 to 100.

2. Every score_breakdown value must be an integer from 0 to 100.

3. Do not invent experience, employers, degrees,
   certifications, dates, skills, achievements, or metrics.

4. If information is not present, say that it is missing.

5. Rewrites must improve wording without changing facts.

6. Do not fabricate numbers.

7. Missing keywords should only be based on the job description.

8. If there is no job description, return:
   "missing_keywords": []
   and "matched_keywords": [].

9. Focus on realistic ATS-readiness rather than claiming
   that any particular ATS will definitely produce this score.

10. Formatting recommendations should focus on ATS parsing:
    standard headings, simple structure, readable text,
    consistent dates, avoiding important information in
    images, excessive tables, text boxes, headers/footers,
    or unusual symbols.

RESUME:
{resume_text[:30000]}

JOB DESCRIPTION:
{
    job_description[:18000]
    if job_description.strip()
    else "No job description was provided."
}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    cleaned = clean_json_response(response.text)

    return json.loads(cleaned)


# ============================================================
# UI HELPERS
# ============================================================

def score_label(score: int) -> str:
    if score >= 85:
        return "Excellent"
    if score >= 70:
        return "Good"
    if score >= 55:
        return "Needs Improvement"
    return "Needs Significant Improvement"


def score_emoji(score: int) -> str:
    if score >= 85:
        return "🟢"
    if score >= 70:
        return "🟡"
    return "🔴"


def safe_list(value: Any) -> List[str]:
    if isinstance(value, list):
        return [str(x) for x in value]
    return []


# ============================================================
# HEADER
# ============================================================

st.markdown(
    '<div class="main-title">📄 AI Resume ATS Analyzer</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="subtitle">'
    "Analyze your resume for ATS-readiness, job-specific "
    "keywords, formatting issues, and improvement opportunities."
    "</div>",
    unsafe_allow_html=True,
)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.header("⚙️ Configuration")

    st.write(f"**Gemini model:** `{MODEL_NAME}`")

    api_key = get_api_key()

    if not api_key:
        entered_key = st.text_input(
            "Gemini API Key",
            type="password",
            placeholder="Paste your Gemini API key",
            help=(
                "For local development only. "
                "For Streamlit Cloud, use Secrets."
            ),
        )

        if entered_key:
            st.session_state.api_key = entered_key.strip()
            api_key = entered_key.strip()

    if api_key:
        st.success("Gemini API key detected.")
    else:
        st.warning("Gemini API key not configured.")

    st.divider()

    st.subheader("Supported files")
    st.write("• PDF")
    st.write("• DOCX")
    st.write("• TXT")

    st.divider()

    st.caption(
        "Privacy note: resume text is sent to Gemini for "
        "AI analysis. Do not upload information you do not "
        "want processed by a third-party AI service."
    )


# ============================================================
# INPUT AREA
# ============================================================

left, right = st.columns([1, 1])

with left:
    uploaded_file = st.file_uploader(
        "📤 Upload your resume",
        type=["pdf", "docx", "txt"],
        help="Upload a text-based PDF, DOCX, or TXT resume.",
    )

with right:
    job_description = st.text_area(
        "💼 Paste the job description",
        height=180,
        placeholder=(
            "Optional but recommended.\n\n"
            "Paste the complete job description here to "
            "calculate job-specific keyword matching."
        ),
    )


# ============================================================
# RESUME PROCESSING
# ============================================================

if uploaded_file is not None:

    if (
        st.session_state.resume_name
        != uploaded_file.name
    ):
        try:
            extracted = extract_resume(uploaded_file)

            st.session_state.resume_text = extracted
            st.session_state.resume_name = uploaded_file.name
            st.session_state.analysis = None

        except Exception as exc:
            st.error(
                f"Unable to read this file: {exc}"
            )
            st.stop()

    resume_text = st.session_state.resume_text

    if not resume_text:
        st.error(
            "No readable text was extracted from the resume. "
            "If your PDF is scanned/image-only, use an OCR-enabled "
            "or text-based PDF."
        )
        st.stop()

    word_count = count_words(resume_text)

    st.success(
        f"Resume loaded successfully: "
        f"**{uploaded_file.name}** · **{word_count} words**"
    )

    with st.expander("👀 Preview extracted resume text"):
        st.text_area(
            "Extracted text",
            resume_text[:15000],
            height=300,
            label_visibility="collapsed",
        )

    # ========================================================
    # LOCAL ATS ANALYSIS
    # ========================================================

    local_result = calculate_local_ats(
        resume_text,
        job_description,
    )

    st.divider()
    st.subheader("📊 Initial ATS Readiness")

    c1, c2, c3, c4 = st.columns(4)

    c1.metric(
        "ATS Readiness",
        f"{local_result['score']}/100",
    )

    c2.metric(
        "Structure",
        f"{local_result['structure_score']}%",
    )

    keyword_score = local_result["keyword"]["score"]

    c3.metric(
        "Keyword Match",
        f"{keyword_score}%"
        if keyword_score is not None
        else "N/A",
    )

    c4.metric(
        "Word Count",
        f"{word_count}",
    )

    st.progress(
        local_result["score"] / 100
    )

    with st.expander("🔍 Local ATS checks"):
        for check, passed in local_result["checks"].items():
            if passed:
                st.write(f"✅ {check}")
            else:
                st.write(f"❌ {check}")

    if local_result["keyword"]["missing"]:
        with st.expander("🔑 Local keyword analysis"):
            matched = local_result["keyword"]["matched"]
            missing = local_result["keyword"]["missing"]

            st.write(
                f"**Matched:** {', '.join(matched[:50])}"
                if matched
                else "**Matched:** None detected"
            )

            st.write(
                f"**Potentially missing:** "
                f"{', '.join(missing[:50])}"
            )

    # ========================================================
    # AI ANALYSIS BUTTON
    # ========================================================

    st.divider()

    analyze_clicked = st.button(
        "🚀 Analyze Resume with Gemini",
        type="primary",
        use_container_width=True,
    )

    if analyze_clicked:

        if not api_key:
            st.error(
                "Gemini API key is required. "
                "Add it in the sidebar or configure "
                "GEMINI_API_KEY in Streamlit Secrets."
            )
            st.stop()

        with st.spinner(
            "Gemini is analyzing your resume..."
        ):
            try:
                result = gemini_analyze(
                    resume_text=resume_text,
                    job_description=job_description,
                    api_key=api_key,
                )

                st.session_state.analysis = result

            except json.JSONDecodeError:
                st.error(
                    "Gemini returned an invalid JSON response. "
                    "Please click Analyze again."
                )

            except Exception as exc:
                st.error(
                    f"Gemini analysis failed: {exc}"
                )


# ============================================================
# DISPLAY AI RESULTS
# ============================================================

analysis = st.session_state.analysis

if analysis:

    st.divider()
    st.header("🤖 AI Resume Analysis")

    try:
        ai_score = int(
            analysis.get(
                "ats_score",
                local_result["score"],
            )
        )
    except (TypeError, ValueError):
        ai_score = local_result["score"]

    ai_score = max(0, min(100, ai_score))

    score_col, summary_col = st.columns([1, 2])

    with score_col:
        st.markdown(
            f"""
            <div class="score-card">
                <div class="small-muted">
                    AI ATS READINESS SCORE
                </div>
                <div class="score-number">
                    {score_emoji(ai_score)} {ai_score}
                </div>
                <div class="score-label">
                    {score_label(ai_score)}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.progress(ai_score / 100)

    with summary_col:
        st.subheader("📝 Overall Assessment")
        st.write(
            analysis.get(
                "summary",
                "No summary was returned.",
            )
        )

    # ========================================================
    # SCORE BREAKDOWN
    # ========================================================

    st.subheader("📈 Score Breakdown")

    breakdown = analysis.get(
        "score_breakdown",
        {},
    )

    b1, b2, b3 = st.columns(3)

    score_items = [
        ("Keyword Match", breakdown.get("keyword_match", 0)),
        ("ATS Formatting", breakdown.get("ats_formatting", 0)),
        ("Skills Relevance", breakdown.get("skills_relevance", 0)),
        ("Experience Relevance", breakdown.get("experience_relevance", 0)),
        ("Achievements", breakdown.get("achievements", 0)),
        ("Resume Structure", breakdown.get("resume_structure", 0)),
    ]

    for index, (label, value) in enumerate(score_items):
        try:
            value = int(value)
        except (TypeError, ValueError):
            value = 0

        value = max(0, min(100, value))

        container = [b1, b2, b3][index % 3]

        with container:
            st.metric(
                label,
                f"{value}/100",
            )
            st.progress(value / 100)

    # ========================================================
    # STRENGTHS / IMPROVEMENTS
    # ========================================================

    strengths = safe_list(
        analysis.get("strengths", [])
    )

    improvements = safe_list(
        analysis.get("top_improvements", [])
    )

    col_a, col_b = st.columns(2)

    with col_a:
        st.subheader("✅ Strengths")

        if strengths:
            for item in strengths:
                st.markdown(f"- {item}")
        else:
            st.info("No strengths returned.")

    with col_b:
        st.subheader("🔧 Top Improvements")

        if improvements:
            for item in improvements:
                st.markdown(f"- {item}")
        else:
            st.info("No improvement recommendations returned.")

    # ========================================================
    # KEYWORDS
    # ========================================================

    missing_keywords = safe_list(
        analysis.get("missing_keywords", [])
    )

    matched_keywords = safe_list(
        analysis.get("matched_keywords", [])
    )

    st.subheader("🔑 Keyword Analysis")

    k1, k2 = st.columns(2)

    with k1:
        st.markdown("**Potentially Missing / Underused**")

        if missing_keywords:
            st.write(", ".join(missing_keywords))
        else:
            st.success(
                "No missing keywords identified."
            )

    with k2:
        st.markdown("**Matched Keywords**")

        if matched_keywords:
            st.write(", ".join(matched_keywords))
        else:
            st.info(
                "No matched keywords returned."
            )

    # ========================================================
    # SECTION FEEDBACK
    # ========================================================

    st.subheader("📌 Section-by-Section Feedback")

    section_feedback = analysis.get(
        "section_feedback",
        {},
    )

    if isinstance(section_feedback, dict):
        for section, feedback in section_feedback.items():
            with st.expander(str(section)):
                st.write(
                    str(feedback)
                )

    # ========================================================
    # REWRITES
    # ========================================================

    rewrites = analysis.get(
        "rewrites",
        [],
    )

    if isinstance(rewrites, list) and rewrites:

        st.subheader("✍️ Suggested Resume Rewrites")

        for index, rewrite in enumerate(rewrites, start=1):

            if not isinstance(rewrite, dict):
                continue

            before = str(
                rewrite.get("before", "")
            )

            after = str(
                rewrite.get("after", "")
            )

            reason = str(
                rewrite.get("reason", "")
            )

            with st.expander(
                f"Rewrite suggestion {index}"
            ):

                st.markdown("**Before**")
                st.info(before)

                st.markdown("**After**")
                st.success(after)

                if reason:
                    st.caption(
                        f"Why: {reason}"
                    )

    # ========================================================
    # FORMATTING / WARNINGS
    # ========================================================

    formatting_tips = safe_list(
        analysis.get("formatting_tips", [])
    )

    warnings = safe_list(
        analysis.get("ats_warnings", [])
    )

    recommended_skills = safe_list(
        analysis.get("recommended_skills", [])
    )

    col_x, col_y = st.columns(2)

    with col_x:
        st.subheader("🎯 ATS Formatting Tips")

        if formatting_tips:
            for item in formatting_tips:
                st.markdown(f"- {item}")
        else:
            st.info("No formatting tips returned.")

    with col_y:
        st.subheader("⚠️ ATS Warnings")

        if warnings:
            for item in warnings:
                st.markdown(f"- {item}")
        else:
            st.success(
                "No major ATS warnings identified."
            )

    # ========================================================
    # RECOMMENDED SKILLS
    # ========================================================

    st.subheader("🧠 Recommended Skills")

    if recommended_skills:
        st.write(
            ", ".join(recommended_skills)
        )
    else:
        st.info(
            "No additional skills were recommended."
        )

    # ========================================================
    # DOWNLOAD REPORT
    # ========================================================

    st.divider()
    st.subheader("📥 Export")

    report = {
        "resume_file": st.session_state.resume_name,
        "local_ats_analysis": local_result,
        "gemini_analysis": analysis,
    }

    report_json = json.dumps(
        report,
        indent=2,
        ensure_ascii=False,
    )

    st.download_button(
        label="⬇️ Download Full Analysis (JSON)",
        data=report_json,
        file_name="resume_ats_analysis.json",
        mime="application/json",
        use_container_width=True,
    )

else:
    if uploaded_file is None:
        st.info(
            "👆 Upload a resume above to start your ATS analysis."
        )

        st.markdown(
            """
            ### How it works

            **1. Upload your resume**  
            PDF, DOCX, or TXT.

            **2. Add a job description**  
            Optional, but strongly recommended for keyword matching.

            **3. Run the local ATS checks**  
            The app checks common resume structure and ATS-readiness signals.

            **4. Analyze with Gemini Flash**  
            Gemini reviews the resume and provides a deeper assessment.

            **5. Improve your resume**  
            Review missing keywords, formatting warnings,
            section feedback, and suggested rewrites.
            """
        )
